"""Los tres exportadores muestran la hora del turno en hora argentina.

2026-09-18, hallazgo S-13 (surgio en la revision de B5-19): la columna Fecha
del PDF era ``item.starts_at.strftime("%Y-%m-%d %H:%M")`` sobre el instante
UTC, sin zona: un turno de las 22:30 ART del dia D salia como ``D+1 01:30``.
El CSV y el Excel exportaban ``isoformat()`` en UTC (``...T01:30:00+00:00``),
correcto para una maquina pero no para el dueno que abre la planilla.
Regla 24: la hora que ve una persona es hora argentina.

Decision para CSV y Excel: hora local legible (``YYYY-MM-DD HH:MM``) con la
zona explicita en el encabezado de la columna, en vez del ISO UTC.
"""

import csv
import io
from datetime import date, time, timedelta

import pytest
from openpyxl import load_workbook

from core.utils import local_to_utc
from modules.reports.exporter import export_to_csv, export_to_excel, export_to_pdf
from modules.reports.schemas import ReportSummaryResponse
from tests.unit.test_report_exporters import _summary
from tests.unit.test_reporte_pdf_columnas import _espiar

DIA = date(2026, 9, 15)
INICIO = "2026-09-15 22:30"
FIN = "2026-09-15 23:00"
ZONA = "America/Argentina/Buenos_Aires"


def _resumen_de_la_noche() -> ReportSummaryResponse:
    resumen = _summary()
    turno = resumen.appointments[0]
    # 22:30 ART del 15 = 01:30Z del 16.
    turno.starts_at = local_to_utc(DIA, time(22, 30))
    turno.ends_at = turno.starts_at + timedelta(minutes=30)
    return resumen


def test_csv_muestra_la_hora_argentina_con_la_zona_en_el_encabezado() -> None:
    contenido = export_to_csv(_resumen_de_la_noche()).decode("utf-8-sig")
    filas = list(csv.reader(io.StringIO(contenido)))
    encabezado = next(f for f in filas if f and f[0] == "public_id")
    turno = next(f for f in filas if f and f[0] == "appt-1")
    fila = dict(zip(encabezado, turno, strict=True))
    assert fila[f"starts_at ({ZONA})"] == INICIO
    assert fila[f"ends_at ({ZONA})"] == FIN


def test_excel_muestra_la_hora_argentina_con_la_zona_en_el_encabezado() -> None:
    libro = load_workbook(io.BytesIO(export_to_excel(_resumen_de_la_noche())))
    hoja = libro["Appointments"]
    filas = [list(fila) for fila in hoja.iter_rows(values_only=True)]
    fila = dict(zip(filas[0], filas[1], strict=True))
    assert fila[f"starts_at ({ZONA})"] == INICIO
    assert fila[f"ends_at ({ZONA})"] == FIN


def test_pdf_muestra_la_hora_argentina(monkeypatch: pytest.MonkeyPatch) -> None:
    trazos = _espiar(monkeypatch)
    export_to_pdf(_resumen_de_la_noche())
    textos = [t.texto for t in trazos]
    assert INICIO in textos
    assert "2026-09-16 01:30" not in textos
